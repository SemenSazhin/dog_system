#!/usr/bin/env python3
"""Краткая презентация для заказчика (10 слайдов, 8-10 минут), строгий стиль, без рекомендаций.
Числа берутся из presentation/cv/stats.json и presentation/cv/detector_compare.json.

python build_customer_deck.py  ->  presentation/customer_report.pptx
"""
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
CV = ROOT / "presentation" / "cv"
OUT = ROOT / "presentation" / "customer_report.pptx"
ST = json.load(open(CV / "stats.json", encoding="utf-8"))["stats"]
DC = json.load(open(CV / "detector_compare.json", encoding="utf-8"))
NAVY, ORANGE, GREY, RED, INK = "#1F4E79", "#E07B39", "#9AA5B1", "#B5483A", "#1B2733"


def A(tag, model, metric="auc", p=3):
    m = ST[tag][model][metric]
    return f"{m['mean']:.{p}f}".replace(".", ","), f"±{m['std']:.{p}f}".replace(".", ",")


def A1(tag, model, metric="auc", p=3):
    a, b = A(tag, model, metric, p)
    return a + b


def c(x):
    return f"{x:.3f}".replace(".", ",")


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


W, H = Inches(13.333), Inches(7.5)


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        r.font.color.rgb = rgb(color)
        p.space_after = Pt(6)
    return tb


def bullets(slide, x, y, w, h, items, size=16):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + it
        r.font.size, r.font.name = Pt(size), "Calibri"
        r.font.color.rgb = rgb(INK)
        p.space_after = Pt(8)
    return tb


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
    else:
        s.line.fill.background()
    s.adjustments[0] = 0.06
    s.shadow.inherit = False
    return s


def table(slide, x, y, rows, col_w, size=13, row_h=0.38, hl=()):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), x, y, Inches(sum(col_w)), Inches(row_h * len(rows)))
    t = shp.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            f = p.runs[0].font
            f.size, f.name = Pt(size), "Calibri"
            f.bold = i == 0 or i in hl
            f.color.rgb = rgb("#FFFFFF") if i == 0 else rgb(INK)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("#DCE8F5" if i in hl else ("#F3F5F7" if i % 2 else "#FFFFFF")))


def pic_fit(slide, path, x, y, max_w, max_h):
    w, h = Image.open(path).size
    k = min(max_w / w, max_h / h)
    slide.shapes.add_picture(str(path), x + int((max_w - w * k) / 2), y, width=int(w * k), height=int(h * k))


prs = Presentation()
prs.slide_width, prs.slide_height = W, H
N = [0]


def new_slide(title, goal=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    N[0] += 1
    add_text(s, Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9), title, size=27, bold=True, color=NAVY)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.12), Inches(1.2), Emu(38000))
    ln.fill.solid()
    ln.fill.fore_color.rgb = rgb(ORANGE)
    ln.line.fill.background()
    add_text(s, Inches(12.2), Inches(7.05), Inches(0.8), Inches(0.35), str(N[0] + 1), size=12, color=GREY, align=PP_ALIGN.RIGHT)
    if goal:
        add_text(s, Inches(0.6), Inches(1.22), Inches(12.1), Inches(0.5), "Цель: " + goal, size=14, color="#5B6B7B")
    return s


def box(slide, text, top=6.25, h=0.7):
    rect(slide, Inches(0.6), Inches(top), Inches(12.1), Inches(h), "#F3F5F7")
    add_text(slide, Inches(0.8), Inches(top + 0.02), Inches(11.7), Inches(h - 0.04), text, size=13, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- график сравнения детекторов
def detector_chart():
    names = ["lstm", "tcn", "transformer", "mamba"]
    fig, axes = plt.subplots(1, 2, figsize=(12.6, 4.1))
    ax = axes[0]
    w = 0.36
    for j, (k, lab, c) in enumerate((("common_old", "прежний детектор", GREY), ("common_new", "yolo26s, imgsz 1280", NAVY))):
        v = [DC["models"][m][k][0] for m in names]
        e = [DC["models"][m][k][1] for m in names]
        ax.bar(np.arange(4) + (j - 0.5) * w, v, w, yerr=e, capsize=3, color=c, label=lab)
        for x, y in zip(np.arange(4) + (j - 0.5) * w, v):
            ax.text(x, y + 0.012, f"{y:.3f}", ha="center", fontsize=9)
    ax.set_xticks(range(4), names)
    ax.set_ylim(0.7, 1.0)
    ax.set_ylabel("ROC-AUC по видео")
    ax.set_title(f"Классификатор на общих {DC['models']['lstm']['common_n']} видео", loc="left", fontsize=12)
    ax.legend(loc="lower right", fontsize=9)
    ax = axes[1]
    ks = [("yolo11n_640", "yolo11n\n640"), ("yolo11n_1280", "yolo11n\n1280"), ("yolo26s_640", "yolo26s\n640"), ("yolo26s_1280", "yolo26s\n1280")]
    v = [DC["frames_120"][k][0] for k, _ in ks]
    ax.bar(range(4), v, color=[GREY, GREY, NAVY, NAVY])
    for i, (k, _) in enumerate(ks):
        ax.text(i, v[i] + 4, f"{v[i]}\nconf {DC['frames_120'][k][1]:.2f}", ha="center", fontsize=9)
    ax.set_xticks(range(4), [l for _, l in ks])
    ax.set_ylim(0, 290)
    ax.set_ylabel("найдено собак на 120 кадрах")
    ax.set_title("Находимость собак (кадры парковых камер)", loc="left", fontsize=12)
    for a_ in axes:
        a_.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    p = CV / "customer_detector.png"
    fig.savefig(p, dpi=170)
    plt.close(fig)
    return p


# ================================================================= 1. титул
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
bg.fill.solid()
bg.fill.fore_color.rgb = rgb(NAVY)
bg.line.fill.background()
add_text(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.0),
         "Детекция дефекации собак по видео", size=42, bold=True, color="#FFFFFF")
add_text(s, Inches(0.9), Inches(3.9), Inches(11.5), Inches(1.2),
         "Отчёт о проделанной работе: эксперименты, результаты, развёртывание", size=22, color="#C9D8E8")
add_text(s, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.5),
         f"{len(ST)} серий экспериментов · 166 видео · оценка по источникам", size=16, color=ORANGE)

# ================================================================= 2. постановка и данные
s = new_slide("Постановка задачи и данные", "зафиксировать, что и на чём измерялось")
bullets(s, Inches(0.6), Inches(1.85), Inches(6.6), Inches(4.4), [
    "Задача: бинарная классификация видео «дефекация / норма» по движению скелета собаки",
    "Данные: 166 видео, 106 класса «норма», 60 класса «дефекация»",
    "64 видео получены с камер парков, и все они относятся к классу «норма»",
    "Метрика: ROC-AUC по видео (одна оценка на ролик), также AP, F1, доля ложных срабатываний на камерах",
    "Протокол: 5 фолдов с разбиением по источникам, несколько сидов, среднее ± стандартное отклонение",
], size=16)
rect(s, Inches(7.6), Inches(1.95), Inches(5.1), Inches(4.1), "#F3F5F7")
add_text(s, Inches(7.85), Inches(2.05), Inches(4.6), Inches(0.5), "Особенность данных", size=17, bold=True, color=NAVY)
add_text(s, Inches(7.85), Inches(2.6), Inches(4.6), Inches(3.4), [
    "Классификатор, которому доступны только разрешение, fps и длительность ролика, даёт ROC-AUC 0,806.",
    "Ролики 1080p составляют 65 из 106 нормальных и 1 из 60 положительных.",
    "Результаты моделей читаются на фоне этой нижней планки.",
], size=15)

# ================================================================= 3. конвейер
s = new_slide("Конвейер обработки", "описать цепочку от кадра до решения")
steps = [("Видеопоток", "YouTube live,\nyt-dlp"), ("Детектор собаки", "YOLO,\nрамка и трекинг"), ("Поза", "20 ключевых\nточек"),
         ("Признаки окна", "нормировка,\nокно T шагов"), ("Классификатор", "LSTM"), ("Решение", "порог, сглаживание,\nподтверждение 2 с")]
x0, bw, gap = 0.6, 1.85, 0.2
for i, (t, d) in enumerate(steps):
    x = x0 + i * (bw + gap)
    rect(s, Inches(x), Inches(2.2), Inches(bw), Inches(1.7), NAVY if i in (1, 4) else "#F3F5F7")
    add_text(s, Inches(x + 0.05), Inches(2.3), Inches(bw - 0.1), Inches(0.5), t, size=15, bold=True,
             color="#FFFFFF" if i in (1, 4) else NAVY, align=PP_ALIGN.CENTER)
    add_text(s, Inches(x + 0.05), Inches(2.95), Inches(bw - 0.1), Inches(0.9), d.split("\n"), size=12,
             color="#FFFFFF" if i in (1, 4) else INK, align=PP_ALIGN.CENTER)
bullets(s, Inches(0.6), Inches(4.3), Inches(12.1), Inches(1.9), [
    "Сервис video-scan: FastAPI в Docker, чтение потоков, сохранение клипов, выгрузка на Яндекс Диск",
    "Ниже оцениваются две изменяемые части: детектор собаки (слайд 8) и классификатор (слайды 5–7)",
    "Модель позы (20 точек) в экспериментах не менялась",
], size=15)

# ================================================================= 4. контроль оценки
s = new_slide("Контроль корректности оценки", "показать, насколько прежние цифры отличались от честных")
table(s, Inches(0.6), Inches(1.95), [
    ["Оценка", "ROC-AUC", "Протокол"],
    ["Первая версия (Colab)", "0,964", "случайное разбиение окон, оценка на окнах"],
    ["Базовая модель lstm", A(*("zoo_nets_noasp", "lstm"))[0], "разбиение по источникам, оценка по видео"],
    ["Лучшая конфигурация lstm", A(*("best_nets", "lstm"))[0], "то же, окно 30 шагов"],
    ["Контроль: только метаданные", "0,806", "разрешение, fps, длительность"],
    ["Контроль: форма рамки", "0,715", "размер и пропорции bbox"],
], [3.9, 1.6, 6.6], size=14, row_h=0.5, hl=(3,))
box(s, "Цифры первой версии завышены утечкой между соседними окнами одного ролика и лёгкостью задачи «камера против стока». "
       "Честная оценка находится в диапазоне 0,85–0,91.", top=5.4, h=0.9)

# ================================================================= 5. архитектуры
s = new_slide("Сравнение архитектур классификатора", "выяснить, зависит ли качество от типа модели")
rows = [["Модель", "ROC-AUC", "AP", "F1 (лучший порог)", "ложные срабатывания на камерах"]]
nets = [("lstm", "LSTM"), ("tcn", "TCN"), ("transformer", "Transformer"), ("mamba", "Mamba"), ("stgcn", "ST-GCN")]
for k, n in nets:
    rows.append([n, A1("zoo_nets_noasp", k), A("zoo_nets_noasp", k, "ap")[0], A("zoo_nets_noasp", k, "f1_best")[0],
                 A("zoo_nets_noasp", k, "cam_fpr_05")[0]])
rows.append(["GBM (признаки по окну)", A1("zoo_gbm_noasp", "gbm"), A("zoo_gbm_noasp", "gbm", "ap")[0],
             A("zoo_gbm_noasp", "gbm", "f1_best")[0], A("zoo_gbm_noasp", "gbm", "cam_fpr_05")[0]])
table(s, Inches(0.6), Inches(1.95), rows, [3.4, 2.0, 1.6, 2.4, 2.7], size=14, row_h=0.5)
add_text(s, Inches(0.6), Inches(5.5), Inches(12.0), Inches(0.4),
         "Базовый протокол: окно 45 шагов, 100 видео, порог доли кадров с позой 0,9. ± — стандартное отклонение по сидам.",
         size=12, color="#5B6B7B")
box(s, "Разница между LSTM, TCN, Transformer и Mamba находится в пределах шума. Механизмы внимания и Mamba выигрыша не дали. "
       "ST-GCN (только точки, без признаков) слабее остальных.", top=6.3, h=0.65)

# ================================================================= 6. группы точек
s = new_slide("Группировка ключевых точек тела", "проверить, повышает ли качество подмножество точек")
pic = CV / "groups_only.png"
if pic.exists():
    pic_fit(s, pic, Inches(0.5), Inches(1.75), Inches(12.3), Inches(3.9))
bullets(s, Inches(0.6), Inches(5.55), Inches(12.1), Inches(0.8), [
    f"Все 20 точек: tcn {A1('grp_all', 'tcn')}; таз+хвост {A1('grp_hips_tail', 'tcn')}; только передние лапы {A1('grp_only_front', 'tcn')} "
    "(с ложными срабатываниями на камерах); хвост и таз по отдельности ниже 0,78",
], size=14)
box(s, "Ни одна группа не превзошла полный набор из 20 точек ни в базовом, ни в лучшем режиме. Важность отдельных групп "
       "нестабильна между режимами: сигнал распределён по всему телу.", top=6.3, h=0.65)

# ================================================================= 7. окно, признаки, лучшая конфигурация
s = new_slide("Окно, признаки и лучшая конфигурация", "найти параметры, дающие максимум на честном протоколе")
rows = [["Конфигурация", "Видео", "lstm", "tcn"]]
for lab, tag, vids in (("T=45, порог 0,9", "zoo_nets_noasp", 100), ("T=30, порог 0,7, простые признаки", "best_nets", 126),
                       ("повтор на новых разбиениях", "best_rep", 126)):
    rows.append([lab, vids, A1(tag, "lstm"), A1(tag, "tcn")])
table(s, Inches(0.6), Inches(1.95), rows, [5.0, 1.2, 2.0, 2.0], size=14, row_h=0.5, hl=(2,))
bullets(s, Inches(0.6), Inches(4.15), Inches(12.1), Inches(2.0), [
    "Лучшая конфигурация: окно 30 шагов (2 с при 15 кадр/с), доля кадров с позой не ниже 0,7, только точки, "
    "без confidence, скоростей и аугментаций",
    "Контроль по метаданным на тех же 126 видео: 0,804. GBM на тех же признаках: " + A1("best_gbm", "gbm"),
    "Окна 60–90 шагов оцениваются на 46–85 видео без части камерных клипов, прямое сравнение с T=30 ограничено",
], size=14)
box(s, "Короткое окно и простые признаки дают максимум: lstm 0,911 и 0,906 на новых разбиениях. "
       "Часть выигрыша относится к изменению состава видео (126 против 100).", top=6.3, h=0.65)

# ================================================================= 8. детектор
s = new_slide("Сравнение детекторов YOLO", "оценить влияние замены детектора собаки на находимость и на классификатор")
pic_fit(s, detector_chart(), Inches(0.5), Inches(1.75), Inches(12.3), Inches(3.5))
m = DC["models"]
bullets(s, Inches(0.6), Inches(5.2), Inches(12.1), Inches(1.0), [
    f"Точки переизвлечены с yolo26s при imgsz 1280: пригодных видео {DC['videos']['new']} против {DC['videos']['old']} (общих {m['lstm']['common_n']})",
    f"Контроль по метаданным на новом наборе видео вырос с {c(DC['ctrl_meta']['old'])} до {c(DC['ctrl_meta']['new'])}, поэтому часть изменений относится к составу видео",
], size=13)
box(s, f"На общих видео lstm {c(m['lstm']['common_old'][0])} и {c(m['lstm']['common_new'][0])} (различий нет), tcn {c(m['tcn']['common_old'][0])} "
       f"и {c(m['tcn']['common_new'][0])}; transformer {c(m['transformer']['common_old'][0])} и {c(m['transformer']['common_new'][0])}, "
       f"mamba {c(m['mamba']['common_old'][0])} и {c(m['mamba']['common_new'][0])}. Находимость собак на 120 кадрах: 153 (прежняя конфигурация сервиса) → 234.",
    top=6.3, h=0.65)

# ================================================================= 9. развёртывание
s = new_slide("Развёртывание сервиса", "зафиксировать текущее состояние production")
table(s, Inches(0.6), Inches(1.95), [
    ["Параметр", "Значение"],
    ["Детектор собаки", "yolo26s, imgsz 1280 (прежде yolo11n, 640)"],
    ["Классификатор", "компактная LSTM, окно 120 кадров, порог 0,8, health-check раз в 6 часов"],
    ["Потоки в анализе", "7: 2 камеры собачьего парка и 5 парковых и уличных камер (YouTube)"],
    ["Хранение клипов", "Яндекс Диск (загрузка работает), локальная копия"],
    ["Нагрузка (CPU, Docker)", "около 360% процессора, память 8,4 ГБ из 15,6 ГБ, стабильна за 10 минут замера"],
    ["Охват оценкой LSTM", f"{DC['service30']['old_scored']} и {DC['service30']['new_scored']} клипов из {DC['service30']['clips']} "
                           "(прежний и новый детектор), окно 120 кадров заполняется не всегда"],
], [3.2, 8.9], size=13, row_h=0.62)
box(s, "Приюты и вольеры были исключены из анализа: дефекации в них не ожидается. Включены парки и городские камеры.", top=6.3, h=0.65)

# ================================================================= 10. ограничения
s = new_slide("Ограничения и статус", "зафиксировать, что результаты не покрывают")
bullets(s, Inches(0.6), Inches(1.9), Inches(12.1), Inches(4.2), [
    "Объём данных: 166 видео (около 22 минут), положительных примеров с камер парков нет, доменный сдвиг камер не измерен",
    "Оценка идёт на уровне ролика, ложные тревоги на непрерывном потоке не измерены",
    "Порог 0,8 в сервисе не калиброван: на клипах «норма» модель давала вероятность выше 0,8 в 4–5 из 7–8 оценённых клипов",
    "Модель в сервисе обучена на точках, полученных прежним детектором, и не переобучалась на yolo26s",
    "Метрика recall при 10% ложных тревог шумная при ~70 нормальных видео, между запусками меняется заметно",
    "Единицы оценки отличаются от первой версии (окна против видео), прямое сравнение цифр некорректно",
], size=16)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
